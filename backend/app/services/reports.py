import io
from datetime import datetime, timezone
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from sqlalchemy.orm import Session
from app.models.enrollment import ApplicantEnrollment, Evidence
from app.models.tracking import TrackingRecord, AuditLog
from app.models.program import PreparatoryCourse, Activity

# Colores institucionales
COLOR_PRIMARY = colors.HexColor("#1E3A8A")   # Azul marino
COLOR_SECONDARY = colors.HexColor("#3B82F6") # Azul claro
COLOR_LIGHT_BG = colors.HexColor("#F3F4F6")
COLOR_RISK_HIGH = colors.HexColor("#DC2626")  # Rojo
COLOR_RISK_MED = colors.HexColor("#D97706")   # Amarillo/Naranja
COLOR_RISK_LOW = colors.HexColor("#16A34A")   # Verde

def generate_applicant_pdf_report(db: Session, enrollment_id: int) -> io.BytesIO:
    """Genera el Reporte Individual del Postulante en formato PDF."""
    enrollment = db.query(ApplicantEnrollment).filter(ApplicantEnrollment.id == enrollment_id).first()
    if not enrollment:
        raise ValueError("Inscripción no encontrada")

    applicant = enrollment.applicant
    program = enrollment.program
    tutor = enrollment.tutor

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=COLOR_PRIMARY,
        alignment=1, # Center
        spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontSize=11,
        textColor=colors.HexColor("#4B5563"),
        alignment=1,
        spaceAfter=15
    )
    section_style = ParagraphStyle(
        "SectionHeader",
        parent=styles["Heading2"],
        fontSize=13,
        leading=16,
        textColor=COLOR_PRIMARY,
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=9,
        leading=12
    )
    bold_body = ParagraphStyle(
        "BoldBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        fontName="Helvetica-Bold"
    )

    elements = []

    # Encabezado
    elements.append(Paragraph("SISTEMA DE GESTIÓN DE CURSOS PREPARATORIOS", title_style))
    elements.append(Paragraph("REPORTE INDIVIDUAL DEL POSTULANTE - SEGUIMIENTO ACADÉMICO", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=COLOR_PRIMARY, spaceAfter=15))

    # 1. Datos Generales y del Programa
    elements.append(Paragraph("1. INFORMACIÓN DEL POSTULANTE Y PROGRAMA", section_style))
    info_data = [
        [
            Paragraph("<b>Nombre Completo:</b>", body_style), Paragraph(applicant.full_name, body_style),
            Paragraph("<b>Programa de Maestría:</b>", body_style), Paragraph(program.name, body_style)
        ],
        [
            Paragraph("<b>Correo Electrónico:</b>", body_style), Paragraph(applicant.email, body_style),
            Paragraph("<b>Convocatoria:</b>", body_style), Paragraph(program.call_for_applications, body_style)
        ],
        [
            Paragraph("<b>Teléfono / Contacto:</b>", body_style), Paragraph(applicant.phone or "No registrado", body_style),
            Paragraph("<b>Tutor Asignado:</b>", body_style), Paragraph(tutor.full_name if tutor else "No asignado", body_style)
        ],
        [
            Paragraph("<b>Estado Actual:</b>", body_style), Paragraph(f"<b>{enrollment.status.value}</b>", body_style),
            Paragraph("<b>Nivel de Riesgo:</b>", body_style), Paragraph(f"<b>{enrollment.risk_level.value}</b>", body_style)
        ],
        [
            Paragraph("<b>Avance General:</b>", body_style), Paragraph(f"<b>{enrollment.completion_percentage}%</b>", body_style),
            Paragraph("<b>Fecha Inscripción:</b>", body_style), Paragraph(enrollment.enrollment_date.strftime("%Y-%m-%d"), body_style)
        ]
    ]

    info_table = Table(info_data, colWidths=[110, 160, 120, 150])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), COLOR_LIGHT_BG),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 10))

    # 2. Cursos Preparatorios y Evidencias Entregadas
    elements.append(Paragraph("2. ACTIVIDADES Y EVIDENCIAS ENTREGADAS", section_style))
    
    evidences_map = {e.activity_id: e for e in enrollment.evidences}
    courses = db.query(PreparatoryCourse).filter(PreparatoryCourse.program_id == program.id).all()

    evidence_table_data = [
        [
            Paragraph("<b>Curso / Actividad</b>", bold_body),
            Paragraph("<b>Ponderación</b>", bold_body),
            Paragraph("<b>Evidencia / Archivo</b>", bold_body),
            Paragraph("<b>Fecha Carga</b>", bold_body),
            Paragraph("<b>Estado</b>", bold_body),
            Paragraph("<b>Calificación</b>", bold_body)
        ]
    ]

    for c in courses:
        for act in c.activities:
            ev = evidences_map.get(act.id)
            ev_status = ev.status.value if ev else "NO_ENTREGADO"
            ev_file = ev.file_name if ev else "-"
            ev_date = ev.upload_date.strftime("%Y-%m-%d") if ev else "-"
            ev_score = str(ev.score) if ev and ev.score is not None else "-"

            evidence_table_data.append([
                Paragraph(f"{c.name} - {act.title}", body_style),
                Paragraph(f"{act.weighting}%", body_style),
                Paragraph(ev_file, body_style),
                Paragraph(ev_date, body_style),
                Paragraph(f"<b>{ev_status}</b>", body_style),
                Paragraph(ev_score, body_style)
            ])

    ev_table = Table(evidence_table_data, colWidths=[180, 65, 115, 65, 75, 40])
    ev_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(ev_table)
    elements.append(Spacer(1, 10))

    # 3. Seguimientos y Observaciones del Tutor
    elements.append(Paragraph("3. HISTORIAL DE SEGUIMIENTOS Y ACCIONES DEL TUTOR", section_style))
    trackings = db.query(TrackingRecord).filter(TrackingRecord.enrollment_id == enrollment.id).order_by(TrackingRecord.contact_date.desc()).all()

    if trackings:
        track_table_data = [
            [
                Paragraph("<b>No. Seguimiento</b>", bold_body),
                Paragraph("<b>Fecha</b>", bold_body),
                Paragraph("<b>Canal</b>", bold_body),
                Paragraph("<b>Resultado / Observaciones</b>", bold_body),
                Paragraph("<b>Compromisos / Próxima Acción</b>", bold_body),
                Paragraph("<b>Estado</b>", bold_body)
            ]
        ]
        for tr in trackings:
            track_table_data.append([
                Paragraph(tr.tracking_number, body_style),
                Paragraph(tr.contact_date.strftime("%Y-%m-%d"), body_style),
                Paragraph(tr.contact_type.value, body_style),
                Paragraph(f"<b>Res:</b> {tr.result}<br/><b>Obs:</b> {tr.observations or '-'}", body_style),
                Paragraph(f"<b>Comp:</b> {tr.commitments or '-'}<br/><b>Prox:</b> {tr.next_action or '-'}", body_style),
                Paragraph(tr.case_status.value, body_style)
            ])
        tr_table = Table(track_table_data, colWidths=[80, 55, 65, 160, 130, 50])
        tr_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#374151")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(tr_table)
    else:
        elements.append(Paragraph("<i>No se han registrado interacciones de tutoría aún.</i>", body_style))

    # Pie de página / Firma
    elements.append(Spacer(1, 20))
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    elements.append(Paragraph(f"Documento generado automáticamente por el Sistema de Cursos Preparatorios el {now_str}.", subtitle_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_consolidated_matrix_excel(db: Session) -> io.BytesIO:
    """Genera la Matriz Consolidada de Seguimiento en formato Excel (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de Seguimiento"

    # Estilos de cabecera
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Título principal
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "MATRIZ CONSOLIDADA DE SEGUIMIENTO A POSTULANTES - MAESTRÍAS"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Nombre Postulante",
        "Programa",
        "Fecha de Inscripción",
        "Avance (%)",
        "Días Pendientes",
        "Tutor Asignado",
        "Núm. Seguimientos",
        "Estado Proceso",
        "Semáforo Riesgo"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    enrollments = db.query(ApplicantEnrollment).all()
    now_date = datetime.now(timezone.utc).date()

    for row_idx, enr in enumerate(enrollments, 4):
        applicant = enr.applicant
        program = enr.program
        tutor = enr.tutor
        
        # Calcular días pendientes hasta el cierre del programa
        days_left = (program.end_date - now_date).days if program and program.end_date else 0
        if days_left < 0:
            days_left = 0

        tracking_count = len(enr.tracking_records)

        row_values = [
            applicant.full_name if applicant else "N/A",
            program.name if program else "N/A",
            enr.enrollment_date.strftime("%Y-%m-%d"),
            f"{enr.completion_percentage}%",
            days_left,
            tutor.full_name if tutor else "Sin asignar",
            tracking_count,
            enr.status.value,
            enr.risk_level.value
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in [3, 4, 5, 7, 8, 9]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Color para el semáforo de riesgo
            if col_idx == 9:
                if val == "ALTO":
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    cell.font = Font(color="DC2626", bold=True)
                elif val == "MEDIO":
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(color="D97706", bold=True)
                else:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    cell.font = Font(color="16A34A", bold=True)

        ws.row_dimensions[row_idx].height = 20

    # Auto-ajuste de ancho de columnas
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(3, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_traceability_excel(db: Session) -> io.BytesIO:
    """Genera el Informe de Trazabilidad de Interacciones en formato Excel (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Trazabilidad de Seguimiento"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "INFORME COMPLETO DE TRAZABILIDAD DE ACCIONES Y TUTORÍA"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "No. Seguimiento",
        "Fecha / Hora",
        "Postulante",
        "Tutor Responsable",
        "Tipo de Contacto",
        "Resultado Obtenido",
        "Compromisos Adquiridos",
        "Respuesta Postulante",
        "Próxima Acción",
        "Estado del Caso"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    trackings = db.query(TrackingRecord).order_by(TrackingRecord.contact_date.desc()).all()

    for row_idx, tr in enumerate(trackings, 4):
        applicant = tr.enrollment.applicant if tr.enrollment else None
        row_values = [
            tr.tracking_number,
            tr.contact_date.strftime("%Y-%m-%d %H:%M"),
            applicant.full_name if applicant else "N/A",
            tr.tutor.full_name if tr.tutor else "N/A",
            tr.contact_type.value,
            tr.result,
            tr.commitments or "-",
            tr.applicant_response or "-",
            tr.next_action or "-",
            tr.case_status.value
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in [1, 2, 5, 10]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

        ws.row_dimensions[row_idx].height = 22

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(3, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
